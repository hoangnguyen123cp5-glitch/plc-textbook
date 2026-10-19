#!/usr/bin/env python3
"""Extract PPTX/PDF textbooks into JSON page dumps under dev/extracts/."""

from __future__ import annotations

import argparse
import sys
import traceback
from pathlib import Path

from common import (
    SOURCES_DIR,
    UNSUPPORTED_SUFFIXES,
    dump_json,
    ensure_dirs,
    extract_json_path,
    first_line,
    iter_source_files,
    source_search_dirs,
)


def iter_shapes_text(shape) -> list[str]:
    texts: list[str] = []
    if getattr(shape, "has_text_frame", False):
        for paragraph in shape.text_frame.paragraphs:
            line = "".join(run.text for run in paragraph.runs) or paragraph.text
            if line and line.strip():
                texts.append(line.strip())
    if getattr(shape, "has_table", False):
        table = shape.table
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                texts.append(" | ".join(cells))
    if getattr(shape, "shapes", None) is not None:
        try:
            for child in shape.shapes:
                texts.extend(iter_shapes_text(child))
        except Exception:
            pass
    return texts


def extract_pptx(path: Path) -> dict:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(str(path))
    pages = []
    for i, slide in enumerate(prs.slides, start=1):
        body: list[str] = []
        title = ""
        if slide.shapes.title and slide.shapes.title.has_text_frame:
            title = slide.shapes.title.text.strip()
        for shape in slide.shapes:
            if shape.has_text_frame and shape == slide.shapes.title:
                continue
            if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.PLACEHOLDER and not shape.has_text_frame:
                continue
            body.extend(iter_shapes_text(shape))
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        text = "\n".join(dict.fromkeys(body))
        pages.append(
            {
                "n": i,
                "title": title or first_line(text),
                "text": text,
                "notes": notes,
            }
        )
    return {
        "file": path.name,
        "path": str(path),
        "kind": "pptx",
        "pageCount": len(pages),
        "pages": pages,
    }


def extract_pdf(path: Path) -> dict:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = []
    for i, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        pages.append(
            {
                "n": i,
                "title": first_line(text),
                "text": text,
                "notes": "",
            }
        )
    return {
        "file": path.name,
        "path": str(path),
        "kind": "pdf",
        "pageCount": len(pages),
        "pages": pages,
    }


def extract_one(path: Path) -> Path:
    suffix = path.suffix.lower()
    if suffix == ".pptx":
        data = extract_pptx(path)
    elif suffix == ".pdf":
        data = extract_pdf(path)
    elif suffix in UNSUPPORTED_SUFFIXES:
        raise ValueError(
            f"{path.name}: {suffix} is not extracted in this version. "
            "Convert .ppt → .pptx (and .doc → PDF) first."
        )
    else:
        raise ValueError(f"{path.name}: unsupported type {suffix}")

    ensure_dirs()
    out = extract_json_path(path)
    dump_json(out, data)
    return out


def collect_inputs(raw: Path | None) -> list[Path]:
    if raw is None:
        dirs = source_search_dirs()
        if not dirs:
            raise FileNotFoundError("No sources/ or Slide/ directory found")
        return iter_source_files(*dirs)
    if raw.is_file() or raw.is_dir():
        return iter_source_files(raw)
    raise FileNotFoundError(f"Not found: {raw}")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Extract PPTX/PDF into dev/extracts/*.json")
    parser.add_argument(
        "path",
        nargs="?",
        default=None,
        help="File or directory (default: sources/ and Slide/)",
    )
    args = parser.parse_args(argv)

    target: Path | None = None
    if args.path:
        target = Path(args.path)
        if not target.is_absolute() and not target.exists():
            alt = Path(__file__).resolve().parents[2] / args.path
            if alt.exists():
                target = alt

    try:
        inputs = collect_inputs(target)
    except FileNotFoundError as err:
        print(err, file=sys.stderr)
        return 1

    if not inputs:
        where = target or SOURCES_DIR
        print(f"No PPTX/PDF files under {where}")
        return 0

    failed = 0
    for path in inputs:
        try:
            out = extract_one(path)
            print(f"OK  {path.name} → {out.relative_to(Path(__file__).resolve().parents[2])}")
        except Exception as err:
            failed += 1
            print(f"ERR {path.name}: {err}", file=sys.stderr)
            if "--debug" in (argv or []):
                traceback.print_exc()
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
